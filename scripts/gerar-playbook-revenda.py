# Gera Documentação/Playbook_Revenda_Portal_BMAX.docx DO ZERO (python-docx).
# Regra do projeto: DOCX é sempre regenerado, nunca editado no XML. Atualize os textos aqui e rode:
#   python -I scripts/gerar-playbook-revenda.py
# Estilo: padrão Boxer (Arial, azul #1D327B, ciano #25BBEE, vermelho #E30613).
# Última revisão: 14/10/2026 — prazos (48 horas úteis sem fim de semana; 60 dias corridos no PCI12a).
import copy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

SAIDA = Path(__file__).resolve().parent.parent.parent / "Documentação" / "Playbook_Revenda_Portal_BMAX.docx"

AZUL, CIANO, VERMELHO = "1D327B", "25BBEE", "E30613"
TEXTO, SEC = "1A202C", "4A5568"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.25)
sec.top_margin = sec.bottom_margin = Cm(1.76)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(10.5)


def run(p, texto, tam=10.5, cor=TEXTO, negrito=False, italico=False):
    r = p.add_run(texto)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(tam)
    r.font.color.rgb = RGBColor.from_string(cor)
    r.bold = negrito or None
    r.italic = italico or None
    return r


def espaco(p, antes=0, depois=0):
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(depois)


def rico(p, partes, tam=10.5):
    """partes: lista de (texto, negrito)."""
    for texto, neg in partes:
        run(p, texto, tam=tam, negrito=neg)


# ── Numeração: lista numerada azul ("1.") e marcadores "•" ─────────────────────────
numbering = doc.part.numbering_part.element
numbering.append(parse_xml(
    f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:multiLevelType w:val="hybridMultilevel"/>'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
    '<w:pPr><w:ind w:left="540" w:hanging="360"/></w:pPr>'
    f'<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:b/><w:color w:val="{AZUL}"/></w:rPr></w:lvl></w:abstractNum>'))
numbering.append(parse_xml(
    f'<w:abstractNum {nsdecls("w")} w:abstractNumId="91"><w:multiLevelType w:val="hybridMultilevel"/>'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
    '<w:pPr><w:ind w:left="540" w:hanging="270"/></w:pPr>'
    f'<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:color w:val="{AZUL}"/></w:rPr></w:lvl></w:abstractNum>'))
# <w:num> precisa vir depois de todas as <w:abstractNum>
numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="91"><w:abstractNumId w:val="91"/></w:num>'))
# abstractNum adicionados no fim ficam depois de <w:num> existentes (o template padrão tem alguns): reordena
nums = [e for e in numbering if e.tag == qn("w:num")]
abstratos = [e for e in numbering if e.tag == qn("w:abstractNum")]
for e in abstratos + nums:
    numbering.remove(e)
for e in abstratos + nums:
    numbering.append(e)


def item(texto_partes, num_id, depois):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'))
    espaco(p, 0, depois)
    rico(p, texto_partes)
    return p


def secao(titulo):
    p = doc.add_paragraph()
    espaco(p, 13, 5)
    p.paragraph_format.keep_with_next = True
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="{CIANO}"/></w:pBdr>'))
    run(p, titulo, tam=13, cor=AZUL, negrito=True)


def faixa(p):
    """Faixa azul de cabeçalho (sombreamento + borda da mesma cor para dar respiro)."""
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="24" w:space="4" w:color="{AZUL}"/>'
        f'<w:left w:val="single" w:sz="24" w:space="4" w:color="{AZUL}"/>'
        f'<w:bottom w:val="single" w:sz="24" w:space="4" w:color="{AZUL}"/>'
        f'<w:right w:val="single" w:sz="24" w:space="4" w:color="{AZUL}"/></w:pBdr>'))
    pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{AZUL}"/>'))


# ── Cabeçalho ───────────────────────────────────────────────────────────────────────
p = doc.add_paragraph(); espaco(p, 6, 0); faixa(p)
run(p, "  PLAYBOOK DA REVENDA", tam=20, cor="FFFFFF", negrito=True)
p = doc.add_paragraph(); espaco(p, 0, 10); faixa(p)
run(p, "  Como usar o Portal BMAX no dia a dia", tam=12, cor="BFD7F5")

p = doc.add_paragraph(); espaco(p, 6, 2)
run(p, "Acesse em ")
p = doc.add_paragraph(); espaco(p, 0, 6)
run(p, "bmax.boxersoldas.com.br", tam=12, cor=AZUL, negrito=True)
run(p, "  •  login pessoal de cada colaborador", cor=SEC)

# ── Rotina em 5 passos ──────────────────────────────────────────────────────────────
secao("Sua rotina em 5 passos")
item([("Abra o painel. ", True), ("Confira os leads novos e os que estão em andamento na sua revenda.", False)], 90, 4.5)
item([("Defina como atender cada lead novo. ", True),
      ("No card, escolha em “Como deseja atender este lead?”: ", False), ("Eu assumo a venda", True), (" ou ", False),
      ("Boxer vende, eu ganho comissão", True),
      (". Você tem até ", False), ("48 horas úteis", True),
      (" (sábado e domingo não contam) e recebe lembrete por e-mail. Sem resposta nesse prazo, o lead segue como "
       "“Boxer vende” e o time Boxer assume o atendimento — você mantém a comissão.", False)], 90, 4.5)
item([("Se você assumiu, atenda e registre o resultado. ", True),
      ("Os dados de contato do cliente passam a aparecer no card. Ao concluir, marque ", False), ("Vendido", True),
      (" (informando o valor real da venda) ou ", False), ("Perdido", True),
      (". Você tem ", False), ("60 dias corridos", True),
      (" a partir do momento em que assumiu; sem resultado nesse prazo o lead é arquivado como “Prazo Vencido”.", False)], 90, 4.5)
item([("Se a Boxer vai vender, acompanhe. ", True),
      ("Você vê o andamento no card e não precisa fazer mais nada: a sua comissão é apurada quando a venda se efetiva.", False)], 90, 4.5)
item([("Trouxe um cliente novo? ", True), ("Use ", False), ("+ Nova Negociação", True),
      (" e preencha os dados, com CNPJ e o equipamento BMAX de interesse.", False)], 90, 4.5)

# ── Regras básicas ──────────────────────────────────────────────────────────────────
secao("Regras básicas")
item([("Só equipamentos BMAX entram no programa ", True),
      ("(máquinas, robôs e laser). Consumíveis, acessórios avulsos e outros produtos não aparecem no portal e não geram cashback.", False)], 91, 3.5)
item([("Escreva o equipamento pelo nome correto, ", True),
      ("como na tabela de preços BMAX (ex.: Hardmig 450, LQ 1250). Nome incompleto ou de outra linha pode deixar o card de fora.", False)], 91, 3.5)
item([("Um cliente (CNPJ) = uma negociação ativa. ", True),
      ("O portal avisa quando já existe um lead aberto para o mesmo CNPJ.", False)], 91, 3.5)
item([("Prazos: ", True),
      ("48 horas úteis para definir o caminho (sábado e domingo não contam; feriados contam) e 60 dias corridos para registrar "
       "Vendido ou Perdido depois de assumir o lead.", False)], 91, 3.5)
item([("Resultado é definitivo. ", True),
      ("Depois de marcar Vendido ou Perdido, o lead é encerrado e não volta a pedir definição de caminho. Marque somente quando for o desfecho real.", False)], 91, 3.5)
item([("Dados corretos desde o início: ", True),
      ("cidade, equipamento e valor. É a base do seu cashback.", False)], 91, 3.5)

# ── Cashback ────────────────────────────────────────────────────────────────────────
secao("Cashback")
item([("Quando gera: ", True), ("nas vendas efetivadas dos leads que aparecem no seu painel.", False)], 91, 3.5)
item([("Onde acompanhar: ", True), ("clique em ", False), ("ver extrato", True),
      (" no card “Cashback Total”. Lá estão o saldo disponível e cada movimentação, com o nome do cliente.", False)], 91, 3.5)
item([("Validade: ", True),
      ("os créditos expiram em 6 meses. O extrato e o e-mail avisam o que está perto de vencer — use antes.", False)], 91, 3.5)
item([("Como usar: ", True), ("no extrato, ", False), ("Solicitar Saque", True),
      (" → informe o valor e o tipo de uso (", False), ("Desconto", True), (" ou ", False), ("Bonificação", True),
      (", com os limites exibidos na tela) → aguarde a aprovação do seu representante.", False)], 91, 3.5)
item([("Depois de aprovado: ", True),
      ("você recebe o cheque com um código (também em PDF). Ele vale por 30 dias: informe o código no campo ", False),
      ("Observação", True), (" do pedido.", False)], 91, 3.5)
item([("O cashback vale para compra de máquinas do portfólio Boxer, não para peças e consumíveis.", False)], 91, 3.5)

# ── Caixa de confidencialidade ──────────────────────────────────────────────────────
p = doc.add_paragraph(); espaco(p, 10, 0)
tabela = doc.add_table(rows=1, cols=1)
tabela.autofit = False
tbl = tabela._tbl
tbl.tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:type="dxa" w:w="9360"/>'))
cel = tabela.rows[0].cells[0]
cel.width = Cm(16.5)
tcPr = cel._tc.get_or_add_tcPr()
tcPr.append(parse_xml(
    f'<w:tcBorders {nsdecls("w")}><w:top w:val="nil"/><w:left w:val="single" w:sz="24" w:color="{VERMELHO}"/>'
    '<w:bottom w:val="nil"/><w:right w:val="nil"/></w:tcBorders>'))
tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F0F4F8"/>'))
tcPr.append(parse_xml(
    f'<w:tcMar {nsdecls("w")}><w:top w:w="140" w:type="dxa"/><w:left w:w="200" w:type="dxa"/>'
    '<w:bottom w:w="140" w:type="dxa"/><w:right w:w="200" w:type="dxa"/></w:tcMar>'))
p = cel.paragraphs[0]; espaco(p, 0, 3); p.paragraph_format.line_spacing = 1.15
run(p, "Confidencialidade", cor=VERMELHO, negrito=True)
for t in ("• Use apenas o seu login e não compartilhe senha. Novo colaborador? Peça o acesso ao seu representante Boxer.",
          "• Dados de clientes, condições e valores do programa são restritos à sua revenda. Não divulgue este material nem prints do portal a terceiros."):
    p = cel.add_paragraph(); espaco(p, 0, 2); p.paragraph_format.line_spacing = 1.15
    run(p, t, tam=10, cor=SEC)

p = doc.add_paragraph(); espaco(p, 7, 0)
run(p, "Dúvidas ou algo estranho no painel? Fale com o seu representante Boxer.", tam=10, cor=SEC, italico=True)

# ── Rodapé com número de página ─────────────────────────────────────────────────────
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, "Uso exclusivo das revendas parceiras BMAX • não compartilhar  |  Página ", tam=8, cor=SEC)
r = run(fp, "", tam=8, cor=SEC)
for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
    if tipo:
        r._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="{tipo}"/>'))
    else:
        r._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve">{texto}</w:instrText>'))

SAIDA.parent.mkdir(exist_ok=True)
doc.save(SAIDA)
print("gerado:", SAIDA)
